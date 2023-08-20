#!/usr/bin/env python

import openpyxl


def BA_type_of_store(shop_type):
    if shop_type == 'OCS':
        return 'COSMETIC STORE'
    elif shop_type == 'OHP' or shop_type == 'OSU' or shop_type == 'OMD' or shop_type == 'OLA' or shop_type == 'OSM':
        return 'GROCERY'
    elif shop_type == 'OMK':
        return 'KIOSK'
    elif shop_type == 'OPS':
        return 'PETROL STATIONS'


path = '/home/layer/Desktop/ba/NewStores.xlsx'


def take_info():
    wb_obj = openpyxl.load_workbook(path)
    worksheet = wb_obj.active

    stores_file = open("stores.txt", "w+", encoding="UTF-8")
    storedtgroup_file = open("storeDTgroup.txt", "w+", encoding="UTF-8")
    storexcodegroup_file = open("storeXgroup.txt", "w+", encoding="UTF-8")
    lbatchstores_file = open("Lbatchstores.txt", "w+", encoding="UTF-8")

    for i in range(2, worksheet.max_row):
        its_done = False
        temp_list = []
        for col in worksheet.iter_cols(1, 14):
            if col[i].value == 'Done':
                its_done = True
                break
            temp_list.append(col[i].value)
        if its_done:
            its_done = False
            continue

        # taking data by postion
        batch = temp_list[2]
        type_of_data = temp_list[3]
        shop_id = temp_list[5]
        xcodegr = temp_list[6]
        shopdescr = temp_list[8]
        retailer = temp_list[9]
        area = temp_list[10]
        surface = temp_list[11]
        region = temp_list[12]
        ac_shopstatus = 'REQUIRED'
        ac_countryid = 'BA'
        ac_languageid = 'BA'
        ac_shoptype = BA_type_of_store(area)
        nc_acv = 100
        nc_xf = 1
        nc_activeflag = 1
        nc_dupitems_flag = 1
        nc_eanxcode_flag = 0
        ac_channelid = 'SCAN'
        nc_dumy_flag = 0
        prefix = '5200'

        if type_of_data == 'SCANNING':
            # fill stores
            stores_list = [prefix + shop_id, shopdescr, ac_shopstatus, xcodegr, ac_countryid, retailer, ac_languageid,
                           ac_shoptype, area, surface, nc_acv, nc_xf, nc_activeflag, nc_dupitems_flag, nc_eanxcode_flag,
                           ac_channelid, nc_dumy_flag, region]

            printer_stores = '\t'.join(str(e) for e in stores_list)
            stores_file.write(printer_stores)
            stores_file.write("\n")

            # fill storedtgroup
            storedtgroup_list = [prefix + shop_id, 'VOLUMETRIC', '0', '1']
            printer_storedtgroup = '\t'.join(str(a) for a in storedtgroup_list)
            storedtgroup_file.write(printer_storedtgroup)
            storedtgroup_file.write("\n")

            # fill storexcodegr
            level_one = [prefix + shop_id, prefix + shop_id, '1', '0', '9999']
            level_two = [prefix + shop_id, xcodegr, '2', '0', '9999']
            level_three = [prefix + shop_id, 'BAXX', '3', '0', '9999']
            level_four = [prefix + shop_id, 'EAN', '4', '0', '9999']

            printer_storeXcode_one = '\t'.join(str(one) for one in level_one)
            printer_storeXcode_two = '\t'.join(str(two) for two in level_two)
            printer_storeXcode_three = '\t'.join(str(three) for three in level_three)
            printer_storeXcode_four = '\t'.join(str(four) for four in level_four)
            storexcodegroup_file.write(printer_storeXcode_one)
            storexcodegroup_file.write("\n")
            storexcodegroup_file.write(printer_storeXcode_two)
            storexcodegroup_file.write("\n")
            storexcodegroup_file.write(printer_storeXcode_three)
            storexcodegroup_file.write("\n")
            storexcodegroup_file.write(printer_storeXcode_four)
            storexcodegroup_file.write("\n")

            # fill lbatchstores

            lbatchstores = [batch, shop_id, prefix + shop_id, '0', '1']
            printer_lbatchstores = '\t'.join(str(l) for l in lbatchstores)
            lbatchstores_file.write(printer_lbatchstores)
            lbatchstores_file.write("\n")




take_info()
print(f'To exist the program press enter')
text = input()